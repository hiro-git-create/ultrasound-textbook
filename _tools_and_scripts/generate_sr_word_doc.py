import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=70, bottom=70, left=90, right=90):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="005691"/>'
        f'<w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def create_sr_doc(output_path):
    doc = docx.Document()
    
    # Page setup - Margins (10mm)
    for section in doc.sections:
        section.top_margin = Inches(0.4)
        section.bottom_margin = Inches(0.4)
        section.left_margin = Inches(0.4)
        section.right_margin = Inches(0.4)
        
    # Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("心エコー検査 DICOM SR連携 推奨項目マスターリファレンス (完全決定版)")
    title_run.font.name = "Meiryo"
    title_run.font.size = Pt(15.5)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0, 70, 130)
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("― 全4弁狭窄（AS/MS/TS/PS）・全4弁逆流（AR/MR/TR/PR）・シャント（Qp/Qs）・PVR・dP/dt・LARS網羅仕様書 ―")
    sub_run.font.name = "Meiryo"
    sub_run.font.size = Pt(9.0)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(70, 70, 70)
    
    # Overview Box
    p_box = doc.add_paragraph()
    p_box.paragraph_format.space_before = Pt(2)
    p_box.paragraph_format.space_after = Pt(4)
    run_box = p_box.add_run(
        "【概要と目的】\n"
        "超音波診断装置で測定した数値をDICOM SRとしてレポートシステム・電子カルテへ自動連携することにより、"
        "① 転記ミスの完全撲滅、② 検査時間の大幅短縮、"
        "③ 【全4弁狭窄症 AS / MS / TS / PS】の連続の式・PHT・圧較差・DVI・Wilkinsスコア、"
        "④ 【全4弁逆流症 AR / MR / TR / PR】のPISA・VC幅・EROA・逆流量・血流逆転、"
        "⑤ 【先進血行動態】Qp/Qs（シャント率）・PVR（肺血管抵抗）・左室 dP/dt・RV-PAカップリング、"
        "⑥ 【最新拡張能 ASE 2025】LARS（左房リザーバーストレイン）・E/e'の完全自動判定を実現します。"
    )
    run_box.font.name = "Meiryo"
    run_box.font.size = Pt(8.0)
    run_box.font.color.rgb = RGBColor(20, 40, 60)
    
    def add_section_header(title_text):
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(7)
        h.paragraph_format.space_after = Pt(2)
        run = h.add_run(title_text)
        run.font.name = "Meiryo"
        run.font.size = Pt(10.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0, 90, 160)
        
    def create_custom_table(headers, rows_data, col_widths=None):
        table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table)
        
        # Header row
        hdr_cells = table.rows[0].cells
        for i, header_text in enumerate(headers):
            hdr_cells[i].text = header_text
            set_cell_background(hdr_cells[i], "005691")
            set_cell_margins(hdr_cells[i], top=70, bottom=70, left=50, right=50)
            p = hdr_cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p.runs:
                r.font.name = "Meiryo"
                r.font.size = Pt(8.2)
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)
                
        # Data rows
        for row_idx, data in enumerate(rows_data):
            row_cells = table.rows[row_idx + 1].cells
            bg_color = "F4F8FA" if row_idx % 2 == 1 else "FFFFFF"
            for col_idx, text in enumerate(data):
                row_cells[col_idx].text = str(text)
                set_cell_background(row_cells[col_idx], bg_color)
                set_cell_margins(row_cells[col_idx], top=45, bottom=45, left=50, right=50)
                p = row_cells[col_idx].paragraphs[0]
                if col_idx in [2, 3]:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif col_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = "Meiryo"
                    r.font.size = Pt(7.8)
                    r.font.color.rgb = RGBColor(40, 40, 40)
                    if col_idx == 1 and any(k in text for k in ["EF", "LAVI", "E/e'", "LARS", "Qp/Qs", "PVR", "EROA", "RVol", "dP/dt", "AVA", "MVA", "TVA", "SVi"]):
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(180, 20, 20)
                        
        if col_widths:
            for row in table.rows:
                for idx, width in enumerate(col_widths):
                    row.cells[idx].width = Inches(width)
        return table

    # 1. Left Ventricle & Hemodynamics
    add_section_header("1. 🫀 左室形態・収縮能 ＆ 特殊血行動態指標（dP/dt）")
    headers_1 = ["大分類", "SR連携推奨項目", "英語表記 / DICOMタグ", "単位", "臨床的意義・システム自動演算"]
    data_1 = [
        ["Mモード / 2D", "左室拡張末期径 / 収縮末期径", "LVDd / LVDs", "mm", "左室径・容量負荷評価・BSA補正値算出"],
        ["Mモード / 2D", "心室中隔壁厚 / 後壁厚", "IVSTd / PWTd", "mm", "壁肥厚・HCM・求心性肥大の評価"],
        ["自動演算", "左室心筋重量係数", "LVMI (LV Mass Index)", "g/m²", "左室肥大 (LVH) 確定診断 (男>115, 女>95)"],
        ["自動演算", "相対的壁厚 / 左室短縮率", "RWT / FS (%FS)", "-, %", "求心性vs遠心性肥大分類 (RWT>0.42), 円周収縮能"],
        ["Simpson法", "左室拡張末期 / 収縮末期容積", "LVEDV / LVESV", "mL", "容量負荷評価・BSA補正 (EDVI, ESVI)"],
        ["Simpson法", "左室駆出率 (最重要)", "LVEF (Biplane EF)", "%", "心機能分類 (HFrEF / HFmrEF / HFpEF)"],
        ["Simpson法", "1回拍出量 / 心拍出量 / SVi", "SV / CO / SVi", "mL, L/min, mL/m²", "有効拍出量・心係数 (CI)・低流量判定 (<35 mL/m²)"],
        ["MR連続波 CW", "左室収縮期圧上昇率 (dP/dt)", "LV dP/dt", "mmHg/s", "MR波形1-3m/s時間より算出。正常>1200, 低下<1000"]
    ]
    create_custom_table(headers_1, data_1, [1.3, 1.8, 1.6, 0.7, 2.1])

    # 2. Diastolic Function & LARS
    add_section_header("2. 🌊 左室拡張能評価（ASE 2025新基準 / LARS・E/e'・LAVI・TR）")
    headers_2 = ["検査手技", "SR連携推奨項目", "英語表記 / DICOMタグ", "単位", "判定基準・自動連携メリット"]
    data_2 = [
        ["左房ストレイン", "左房リザーバーストレイン (新基準)", "LARS (LA Res. Strain)", "%", "充満圧上昇確定 (重度低下 < 18%, 軽度低下 < 24%, 正常 > 39%)"],
        ["左房ストレイン", "左房コンジット / ポンプストレイン", "LACS / LAAS", "%", "受動的導管機能 / 心房能動的収縮機能の評価"],
        ["TVI (流入血流)", "E波最高血流速度 / A波最高血流速度", "MV E vel / MV A vel", "cm/s", "拡張早期 / 心房収縮期流入速度 (AF時はA波欠損)"],
        ["自動演算", "E/A比 / E波減速時間 (DT)", "E/A ratio / MV DT", "-, ms", "弛緩障害型 (<0.8), 拘束型 (>2.0, DT<160ms)"],
        ["組織ドプラ TDI", "中隔側 / 側壁側 e' 速度", "Septal e' / Lateral e'", "cm/s", "局所弛緩能低下 (中隔 < 7.0, 側壁 < 10.0 cm/s)"],
        ["自動演算", "平均 E/e' 比 (最重要)", "Average E/e'", "-", "左房圧・充満圧指標 (> 14 で上昇, AF時は ≧ 11)"],
        ["Biplane容積", "左房容積係数 (最重要)", "LAVI", "mL/m²", "慢性左房圧上昇 (> 34 mL/m² で拡大陽性)"],
        ["連続波ドプラ CW", "三尖弁逆流最高流速", "TR Vmax", "m/s", "肺動脈圧上昇スクリーニング (> 2.8 m/s で陽性)"],
        ["肺静脈血流 PW", "S/D比 ＆ Ar-A持続時間差", "PV S/D / (Ar dur - A dur)", "-, ms", "S < D: 左房圧上昇 / Ar-A ≧ 30ms: LVEDP上昇"]
    ]
    create_custom_table(headers_2, data_2, [1.3, 1.8, 1.6, 0.7, 2.1])

    # 3. Complete Valvular Stenosis Quantification (AS, MS, TS, PS)
    add_section_header("3. 🔒 全弁膜狭窄症の完全定量・重症度評価（AS / MS / TS / PS）")
    p_sten = doc.add_paragraph()
    r_sten = p_sten.add_run(
        "【弁狭窄評価の核心】\n"
        "・AS: 連続の式 AVA < 1.0 cm² (AVAi < 0.6 cm²/m²), Vmax ≧ 4.0 m/s, Mean PG ≧ 40 mmHg, DVI < 0.25 (LFLG判定 SVi < 35)\n"
        "・MS: PHT法 MVA ≦ 1.5 cm² (重症 ≦ 1.0 cm²), Mean PG ≧ 10 mmHg, PHT ≧ 220 ms, Wilkinsスコア (PTMC適応 ≦ 8点)\n"
        "・TS: Mean PG ≧ 5 mmHg, PHT ≧ 190 ms, TVA ≦ 1.0 cm²  |  PS: Vmax > 4.0 m/s, Max PG ≧ 64 mmHg (重症)"
    )
    r_sten.font.name = "Meiryo"
    r_sten.font.size = Pt(7.8)
    r_sten.font.bold = True
    r_sten.font.color.rgb = RGBColor(160, 50, 0)

    headers_3 = ["対象弁狭窄", "SR連携推奨項目", "英語表記 / DICOMタグ", "単位", "重症度判定カットオフ値 (ASE/ESC基準)"]
    data_3 = [
        ["大動脈弁狭窄 AS", "大動脈弁最高血流速度", "AV Vmax (Peak Vel)", "m/s", "重症: ≧ 4.0 m/s (中等症: 3.0〜4.0 m/s)"],
        ["大動脈弁狭窄 AS", "大動脈弁平均圧較差 / 最大PG", "AV Mean PG / AV Max PG", "mmHg", "重症: Mean PG ≧ 40 mmHg (Max PG ≧ 64 mmHg)"],
        ["大動脈弁狭窄 AS", "大動脈弁口面積 (連続の式)", "AVA (Continuity Eq)", "cm²", "重症: < 1.0 cm² (体表面積補正 AVAi < 0.6 cm²/m²)"],
        ["大動脈弁狭窄 AS", "無次元速度指数 (DVI)", "DVI (LVOT VTI / AV VTI)", "-", "重症: < 0.25 (低流量低圧較差 LF-LG AS で極めて有用)"],
        ["大動脈弁狭窄 AS", "大動脈弁加速時間 / 駆出時間比", "AV AcT / (AcT / ET)", "ms, -", "重症: AcT ≧ 100 ms, AcT/ET > 0.35 (弁開放遅延)"],
        ["大動脈弁狭窄 AS", "低流量判定 (1回拍出量係数)", "SVi (Stroke Volume Index)", "mL/m²", "SVi < 35 mL/m² で Low-flow severe AS (奇異性LFLG)"],
        ["僧帽弁狭窄 MS", "僧帽弁平均圧較差 / 最大流速", "MV Mean PG / MV Peak E", "mmHg, m/s", "重症: Mean PG ≧ 10 mmHg (中等症: 5〜10 mmHg)"],
        ["僧帽弁狭窄 MS", "僧帽弁圧半減時間 (PHT)", "MV PHT", "ms", "重症: ≧ 220 ms (MVA = 220 / PHT)"],
        ["僧帽弁狭窄 MS", "僧帽弁口面積 (PHT法 / トレース)", "MVA (PHT) / MVA (Plan)", "cm²", "重症: ≦ 1.0 cm² (臨床的有意狭窄: ≦ 1.5 cm²)"],
        ["僧帽弁狭窄 MS", "連続の式僧帽弁口面積", "MVA (Continuity Eq)", "cm²", "AR合併等でPHTの信頼性が低下した際の確定評価"],
        ["僧帽弁狭窄 MS", "Wilkinsエコースコア (4項目)", "Wilkins Score (Total)", "点", "弁尖肥厚・可動性・石灰化・下部病変 (PTMC適応 ≦ 8点)"],
        ["三尖弁狭窄 TS", "三尖弁平均圧較差 / 圧半減時間", "TV Mean PG / TV PHT", "mmHg, ms", "重症: Mean PG ≧ 5 mmHg / PHT ≧ 190 ms"],
        ["三尖弁狭窄 TS", "三尖弁口面積 (PHT法 / 連続の式)", "TVA (PHT) / TVA (Cont)", "cm²", "重症: ≦ 1.0 cm² (TVA = 190 / TV PHT)"],
        ["肺動脈弁狭窄 PS", "肺動脈弁最高血流速度 / 最大PG", "PV Vmax / PV Max PG", "m/s, mmHg", "重症: Vmax > 4.0 m/s / Max PG ≧ 64 mmHg (中等症: 36〜64)"],
        ["肺動脈弁狭窄 PS", "肺動脈弁平均圧較差", "PV Mean PG", "mmHg", "重症: ≧ 35 mmHg"]
    ]
    create_custom_table(headers_3, data_3, [1.3, 1.8, 1.6, 0.7, 2.1])

    # 4. Complete Valvular Regurgitation Quantification (AR, MR, TR, PR)
    add_section_header("4. 🎯 全弁膜逆流症の完全定量・半定量評価（AR / MR / TR / PR）")
    headers_4 = ["対象弁逆流", "SR連携推奨項目", "英語表記 / DICOMタグ", "単位", "重症度判定カットオフ値 (ASE/ESC基準)"]
    data_4 = [
        ["大動脈弁逆流 AR", "Vena Contracta 幅 (VC)", "AR VC width", "mm", "重症: > 6.0 mm (軽症: < 3.0 mm)"],
        ["大動脈弁逆流 AR", "圧半減時間 (PHT)", "AR PHT", "ms", "重症: < 200 ms (軽症: > 500 ms)"],
        ["大動脈弁逆流 AR", "有効逆流弁口面積 (EROA)", "AR EROA", "cm²", "重症: ≧ 0.30 cm² (軽症: < 0.10 cm²)"],
        ["大動脈弁逆流 AR", "逆流量 (Regurgitant Volume)", "AR RVol", "mL", "重症: ≧ 60 mL (軽症: < 30 mL)"],
        ["大動脈弁逆流 AR", "下行大動脈拡張期逆流終末速度", "AR Holodiastolic flow (EDV)", "cm/s", "重症: 下行大動脈全拡張期逆流 ＆ EDV > 20 cm/s"],
        ["僧帽弁逆流 MR", "Vena Contracta 幅 (VC)", "MR VC width", "mm", "重症: ≧ 7.0 mm (軽症: < 3.0 mm)"],
        ["僧帽弁逆流 MR", "PISA半径 / アライアンス速度", "PISA Radius / Aliasing Vel", "mm, cm/s", "定量的逆流評価 (PISA法) の必須元データ"],
        ["僧帽弁逆流 MR", "有効逆流弁口面積 (EROA)", "MR EROA", "cm²", "重症: ≧ 0.40 cm² (二次性MRでは ≧ 0.20 cm²)"],
        ["僧帽弁逆流 MR", "逆流量 (Regurgitant Volume)", "MR RVol", "mL", "重症: ≧ 60 mL (二次性MRでは ≧ 30 mL)"],
        ["僧帽弁逆流 MR", "肺静脈逆流波 (収縮期逆流)", "PV Systolic flow reversal", "-", "重症: 肺静脈血流で収縮期逆流 (S波の陰転化)"],
        ["三尖弁逆流 TR", "三尖弁逆流最高流速 / 最大PG", "TR Vmax / TR max PG", "m/s, mmHg", "肺動脈圧推定の基幹 (Vmax > 2.8 m/s でPH疑い)"],
        ["三尖弁逆流 TR", "TR Vena Contracta 幅 (VC)", "TR VC width", "mm", "重症: ≧ 7.0 mm (Massive 14-20, Torrential ≧21)"],
        ["三尖弁逆流 TR", "TR PISA EROA / 逆流量", "TR EROA / TR RVol", "cm², mL", "重症: EROA ≧ 0.40 cm² / RVol ≧ 45 mL"],
        ["三尖弁逆流 TR", "肝静脈収縮期逆流波", "Hepatic vein flow reversal", "-", "重症: 肝静脈波形での収縮期逆流 (Blunting/Reversal)"],
        ["肺動脈弁逆流 PR", "PR peak vel / end-diastolic vel", "PR peak vel / PRed vel", "m/s", "平均肺動脈圧 (mPAP) ＆ 拡張期圧 (PADP) 推定"],
        ["肺動脈弁逆流 PR", "PR 圧半減時間 (PHT)", "PR PHT", "ms", "重症: < 100 ms で急峻な減衰 (Severe PR)"],
        ["肺動脈弁逆流 PR", "PR Index (持続時間比)", "PR Index (PR dur / Diastole)", "-", "重症: < 0.77 (拡張期の早期に血流途絶)"]
    ]
    create_custom_table(headers_4, data_4, [1.3, 1.8, 1.6, 0.7, 2.1])

    # 5. Advanced Hemodynamics: Qp/Qs, PVR, RV-PA Coupling
    add_section_header("5. 🫁 先進血行動態指標（Qp/Qs・PVR・RV-PAカップリング）")
    headers_5 = ["演算領域", "SR連携推奨項目", "英語表記 / DICOMタグ", "単位", "計算ロジック・臨床判断基準"]
    data_5 = [
        ["シャント率 (Qp/Qs)", "肺体血流比 (Qp/Qs)", "Qp/Qs ratio", "-", "ASD/VSD/PDAの評価。正常 1.0, > 1.5 で閉鎖適応"],
        ["Qp/Qs 基礎データ", "右室流出道径 / RVOT VTI", "RVOT diam / RVOT VTI", "mm, cm", "Qp (肺血流量) 算出のための必須計測"],
        ["Qp/Qs 基礎データ", "左室流出道径 / LVOT VTI", "LVOT diam / LVOT VTI", "mm, cm", "Qs (体血流量) 算出のための必須計測"],
        ["肺血管抵抗 (PVR)", "肺血管抵抗 (Abbas推定式)", "PVR (Wood Units)", "Wood U", "正常 < 2.0, > 3.0 で前毛細管性肺高血圧 (毛細血管病変)"],
        ["右室PA連関", "TAPSE / PASP 比 (カップリング)", "TAPSE/PASP ratio", "mm/mmHg", "右室後負荷不整合の指標。正常 > 0.55, 予後不良 < 0.36"],
        ["右室心筋機能", "右室 Tei Index (MPI)", "RV Tei Index (MPI)", "-", "組織ドプラで > 0.54 (パルスで > 0.43) で機能低下"]
    ]
    create_custom_table(headers_5, data_5, [1.3, 1.8, 1.6, 0.7, 2.1])

    # 6. Right Heart, IVC, Aorta
    add_section_header("6. 📏 右室形態・下大静脈・大動脈基部計測")
    headers_6 = ["評価領域", "SR連携推奨項目", "英語表記 / DICOMタグ", "単位", "判定基準・カットオフ"]
    data_6 = [
        ["右室収縮能", "三尖弁輪収縮期移動距離", "TAPSE", "mm", "< 17 mm で右室収縮能低下"],
        ["右室収縮能", "組織ドプラ三尖弁輪収縮速度", "RV s' (TDI S')", "cm/s", "< 9.5 cm/s で右室収縮能低下"],
        ["右室収縮能", "右室面積変化率", "RV FAC", "%", "< 35 % で右室機能低下"],
        ["下大静脈 IVC", "下大静脈最大径 / 虚脱率", "IVCd max / IVC Collapse", "mm, %", "> 21 mm ＆ 虚脱率 < 50% で右房圧上昇 (RAP 15mmHg)"],
        ["自動演算", "推定右房圧 / 推定肺動脈収縮期圧", "RAP / PASP (TR-PG + RAP)", "mmHg", "RAP: 3/8/15 mmHg, PASP > 35〜40 mmHg で肺高血圧"],
        ["大動脈基部", "弁輪 / Valsalva / STJ / 上行径", "Ao Annulus/Sinus/STJ/Asc", "mm", "Valsalva / 上行 > 40 mm で拡大 (≧ 50mm 手術検討)"],
        ["心膜腔", "心嚢液深度 (拡張末期)", "Pericardial Effusion Depth", "mm", "少量 <10mm, 中等量 10-20mm, 大量 >20mm"]
    ]
    create_custom_table(headers_6, data_6, [1.3, 1.8, 1.6, 0.7, 2.1])

    # Operational & System Engineering
    add_section_header("💡 現場でのSR連携 運用・システム設計チェックポイント")
    p_tips = doc.add_paragraph()
    r_tips = p_tips.add_run(
        "1. 計測ラベルの選択厳守: フリーキャリパーではなく、必ず装置内蔵の専用ラベル（例: Ao Diam, MV E, Sep e', LARS, RVOT diam, PISA, Wilkins等）を選択して計測すること。\n"
        "2. 弁狭窄の連続の式ペアリング: LVOT径・LVOT VTI と AV/MV/TV/PV の各弁VTIが揃って初めて狭窄弁口面積（AVA/MVA/TVA）が完全自動算出される。\n"
        "3. シャント率 (Qp/Qs) ＆ PVRの自動計算: RVOTとLVOTの径・VTI、および TR Vmax からQp/QsとPVR（Wood単位）を完全自動演算。\n"
        "4. 全4弁狭窄・全4弁逆流の完全網羅: AS/MS/TS/PS および AR/MR/TR/PR の全重症度基準値を網羅し、弁膜症レポート作成を自動化する。\n"
        "5. 単位系スケーリングの整合性: エコー機側の出力単位（cm/s ⇄ m/s、mL ⇄ L）とレポートシステム側の受信単位の整合性を結合テストで必ず照合すること。\n"
        "6. 複数計測の代表値採用ルール: 不整脈や連続波ドプラ等で複数回計測した場合、レポート側で「平均値（Average）」を採用する設定に固定することを推奨する。"
    )
    r_tips.font.name = "Meiryo"
    r_tips.font.size = Pt(7.8)
    r_tips.font.color.rgb = RGBColor(30, 30, 30)

    doc.save(output_path)
    print(f"Successfully updated master complete docx at {output_path}")

if __name__ == "__main__":
    local_path = r"c:\Antigravity\超音波検査\心エコーデジタル教科書\心エコー_DICOM_SR連携_推奨項目マスターリファレンス.docx"
    create_sr_doc(local_path)
